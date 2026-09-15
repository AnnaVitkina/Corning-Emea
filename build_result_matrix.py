"""Build result matrix from a processing/ DataFrame and save to output/.

Customize mappings, cost columns, and equipment rules in this module.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from excel_io import load_to_dataframe
from paths import OUTPUT_DIR, PROCESSING_DIR

SUPPORTED_EXTENSIONS = {".xlsx", ".xls", ".csv", ".tsv"}

# Equipment Type Clean -> size suffix(es) used in cost headers
EQUIPMENT_SIZE_TARGETS: dict[str, list[str]] = {
    "20' Standard Dry": ["20FT"],
    "40' Standard Dry/40' High Cube Dry": ["40FT", "40HC"],
    "40' Standard Dry": ["40FT"],
    "40' High Cube Dry": ["40HC"],
}

# Mode words in Origin Mode / Dest Mode -> mode family for cost headers
MODE_ALIASES: dict[str, str] = {
    "truck": "Truck",
    "motor": "Truck",
    "barge": "Barge",
    "rail": "Rail",
}
MODE_ORDER = ["Truck", "Barge", "Rail"]
SIZE_ORDER = ["20FT", "40FT", "40HC"]
DEFAULT_MODE = "Truck"
COMBINED_SERVICE_LABEL = "Standard/Pre-delivered"

# Candidate source columns for ICL "Multi stop"
MULTI_STOP_COLUMNS = [
    "Multi stop",
    "Multi-stop",
    "Multistop",
    "Multi Stop",
    "Accessorial Flex Field 2",
]

# Destination countries treated as EMEA import destinations
EMEA_COUNTRY_CODES = {
    "AL", "AD", "AM", "AT", "AZ", "BA", "BE", "BG", "BY", "CH", "CY", "CZ", "DE",
    "DK", "EE", "ES", "FI", "FR", "GB", "GE", "GR", "HR", "HU", "IE", "IS", "IT",
    "KZ", "LI", "LT", "LU", "LV", "MC", "MD", "ME", "MK", "MT", "NL", "NO", "PL",
    "PT", "RO", "RS", "RU", "SE", "SI", "SK", "SM", "TR", "UA", "UK", "VA", "XK",
}

# Optional shorter / renamed names for cost header row 1
COST_DISPLAY_NAMES: dict[str, str] = {
    "All-In Origin Charges (Except Inland Trucking/Rail)": "All-In Origin Charges",
    "All-In Destination Charges (Except Inland Trucking/Rail)": "All-In Destination Charges",
    "Freight Rates": "Transport cost",
    "Mandatory Misc Surcharge Amount": "Terminal Handling Charges",
    "OTHC/DTHC - Terminal Handling Charges": "Terminal Handling Charges",
    "Europe Emissions Surcharge": "ETS Fee",
    "ETS Fee": "ETS Fee",
    "ETS Fee p. Container": "ETS Fee p. Container",
    "Empty Container Drop Off p. Cont.": "Empty Container Drop Off",
    "Multi stop": "Multi stop",
    "Multi-stop": "Multi stop",
    "Multistop": "Multi stop",
    "Multi Stop": "Multi stop",
    "Accessorial Flex Field 2": "Multi stop",
    "B/L Fee per Shipment": "B/L Fee",
    "Export Clearance Fee per Shipment": "Export Clearance Fee",
    "Import Clearance Fee per Shipment": "Import Clearance Fee",
    "Doc Turnover Fees per Shipment": "Doc Turnover Fee",
}

# Costs split by equipment type (Truck/Barge/Rail x size)
EQUIPMENT_COST_COLUMNS = [
    "All-In Origin Charges (Except Inland Trucking/Rail)",
    "Origin Inland Trucking/Rail",
    "Freight Rates",
    "Destination Inland Trucking/Rail",
    "All-In Destination Charges (Except Inland Trucking/Rail)",
    "BAF/Fuel",
    "Min. Charge per Shipment",
    "VGM Fee/Container",
    "Empty Container Drop Off p. Cont.",
    "Driver Wait Time/Hour",
    "Detention at Origin per container per day",
    "Demurrage at Origin per container per day",
    "Detention at Destination per container per day",
    "Demurrage at Destination per container per day",
    "Chassis Rental per container per day",
    "Management Fee",
    "Accessorial Flex Field 2",
    "Accessorial Flex Field 3",
    "Accessorial Flex Field 4",
    "Accessorial Flex Field 5",
    "Europe Emissions Surcharge",
    "ETS Fee",
    "ETS Fee p. Container",
    "Interim Disruption Surcharge Amount",
    "Mandatory Misc Surcharge Amount",
    "OTHC/DTHC - Terminal Handling Charges",
]

# Shared costs: one block at the end, applied to all equipment types (not split)
# Always included when present — including when values are 0
SHARED_COST_COLUMNS = [
    "B/L Fee per Shipment",
    "Export Clearance Fee per Shipment",
    "Import Clearance Fee per Shipment",
    "Doc Turnover Fees per Shipment",
]

SHARED_RATE_BY = "Rate by: All equipment"


@dataclass
class RateCardProfile:
    """Filename-based conversion rules."""

    name: str
    use_origin_dest_modes: bool = False
    allowed_modes: tuple[str, ...] = ("Truck", "Barge", "Rail")
    allowed_sizes: tuple[str, ...] | None = None  # None = from equipment type
    force_targets: tuple[str, ...] | None = None  # e.g. PISA -> Truck/40FT only
    mirror_40ft_40hc: bool = False
    prefer_dest_mode: bool = False  # ICL: mode from destination seaport
    service_label: str | None = None
    usd_only: bool = False
    import_only: bool = False
    equipment_cost_whitelist: tuple[str, ...] | None = None
    include_shared_costs: bool = True


DEFAULT_PROFILE = RateCardProfile(name="default")

SCHENKER_PROFILE = RateCardProfile(
    name="schenker",
    use_origin_dest_modes=True,
    allowed_modes=("Truck", "Barge"),
    mirror_40ft_40hc=True,
)

ICL_PROFILE = RateCardProfile(
    name="icl",
    use_origin_dest_modes=True,
    allowed_modes=("Truck", "Barge"),
    allowed_sizes=("40FT", "40HC"),
    mirror_40ft_40hc=True,
    prefer_dest_mode=True,
    service_label=COMBINED_SERVICE_LABEL,
    usd_only=True,
    import_only=True,
    equipment_cost_whitelist=("Freight Rates", *MULTI_STOP_COLUMNS),
    include_shared_costs=False,
)

PISA_PROFILE = RateCardProfile(
    name="pisa",
    force_targets=("Truck/40FT",),
)

ROHLIG_PROFILE = RateCardProfile(
    name="rohlig",
    use_origin_dest_modes=True,
    allowed_modes=("Truck", "Barge", "Rail"),
    allowed_sizes=("20FT", "40FT", "40HC"),
    mirror_40ft_40hc=True,
    service_label=COMBINED_SERVICE_LABEL,
)


def detect_rate_card_profile(source: Path | str | None) -> RateCardProfile:
    name = Path(source).name.lower() if source is not None else ""
    if "rohlig" in name:
        return ROHLIG_PROFILE
    if "schenker" in name:
        return SCHENKER_PROFILE
    if "icl" in name:
        return ICL_PROFILE
    if "pisa" in name:
        return PISA_PROFILE
    return DEFAULT_PROFILE


HEADER_FILL = PatternFill("solid", fgColor="D9D9E8")
ZEBRA_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)
HEADER_FONT = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")


@dataclass
class CostBlock:
    """One cost group = 2 Excel columns (Currency/% + value basis)."""

    title: str
    rate_by: str
    basis: str
    currency: pd.Series
    values: pd.Series
    left_header: str = "Currency"


@dataclass
class ResultMatrix:
    shipment: pd.DataFrame
    costs: list[CostBlock]

    @property
    def shape(self) -> tuple[int, int]:
        return (len(self.shipment), len(self.shipment.columns) + 2 * len(self.costs))


def list_processing_files() -> list[Path]:
    if not PROCESSING_DIR.exists():
        raise FileNotFoundError(f"Processing folder not found: {PROCESSING_DIR}")

    return sorted(
        p
        for p in PROCESSING_DIR.iterdir()
        if p.is_file()
        and not p.name.startswith("~$")
        and p.suffix.lower() in SUPPORTED_EXTENSIONS
    )


def ask_user_to_select(files: list[Path]) -> Path:
    print("\nFiles available in processing/:\n")
    for i, path in enumerate(files, start=1):
        print(f"  {i}. {path.name}")

    while True:
        choice = input(f"\nSelect a file to convert to result matrix [1-{len(files)}]: ").strip()
        if not choice.isdigit():
            print("Please enter a number.")
            continue
        index = int(choice)
        if 1 <= index <= len(files):
            return files[index - 1]
        print(f"Please choose a number between 1 and {len(files)}.")


def format_date_ddmmyyyy(value: object) -> str | object:
    if pd.isna(value):
        return value
    ts = pd.to_datetime(value, errors="coerce")
    if pd.isna(ts):
        return value
    return ts.strftime("%d.%m.%Y")


def map_carrier_name(carrier: object, destination_country: object) -> object:
    if pd.isna(carrier):
        return carrier
    carrier_text = str(carrier).strip().upper().replace(" ", "")
    # DSV/SCHENKER (and spelling variants) -> "DSV" + destination country
    if "DSV" in carrier_text and "SCHENKER" in carrier_text.replace("SHENCKER", "SCHENKER"):
        dest = "" if pd.isna(destination_country) else str(destination_country).strip()
        return f"DSV {dest}".strip()
    return carrier


def clean_text(value: object) -> object:
    """Normalize shipment text/codes: blank for missing, strip floats like 6012.0."""
    if value is None or pd.isna(value):
        return pd.NA
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).strip()
    text = str(value).strip()
    if text == "" or text.lower() in {"nan", "none", "nat"}:
        return pd.NA
    if text.endswith(".0"):
        head = text[:-2]
        if head.isdigit() or (head.startswith("-") and head[1:].isdigit()):
            return head
    return text


def optional_series(df: pd.DataFrame, *candidates: str) -> pd.Series:
    """Return first matching source column, cleaned; else an empty series."""
    for name in candidates:
        if name in df.columns:
            return df[name].map(clean_text)
    return pd.Series([pd.NA] * len(df), dtype="object")


def rate_basis(cost_name: str) -> str:
    """Return 'Flat' or 'p/unit' for the cost header third row."""
    lower = cost_name.lower()
    if "per shipment" in lower:
        return "Flat"
    return "p/unit"


def cost_title(cost_name: str, equipment: str | None = None) -> str:
    display = COST_DISPLAY_NAMES.get(cost_name, cost_name)
    if equipment is None:
        return display
    return f"{display} ({equipment})"


def rate_by_label(equipment: str) -> str:
    return f"Rate by: {equipment}"


def is_rohlig_rate_card(source: Path | str | None) -> bool:
    return detect_rate_card_profile(source).name == "rohlig"


def apply_service_label(df: pd.DataFrame, label: str) -> pd.DataFrame:
    out = df.copy()
    out["Service Level Clean"] = label
    return out


def filter_profile_rows(df: pd.DataFrame, profile: RateCardProfile) -> pd.DataFrame:
    """Apply USD / import lane filters for profiles that require them."""
    out = df
    if profile.usd_only and "Currency" in out.columns:
        out = out[out["Currency"].astype(str).str.strip().str.upper() == "USD"]
    if profile.import_only and "Destination Country Clean" in out.columns:
        dest = out["Destination Country Clean"].astype(str).str.strip().str.upper()
        out = out[dest.isin(EMEA_COUNTRY_CODES)]
    return out.reset_index(drop=True)


def parse_modes_from_cell(value: object) -> list[str]:
    """Extract Truck/Barge/Rail from a mode cell (supports Rail/Truck, etc.)."""
    if value is None or pd.isna(value):
        return []
    text = str(value).strip()
    if not text:
        return []

    found: list[str] = []
    parts = []
    for chunk in text.replace("&", "/").replace(",", "/").replace(";", "/").split("/"):
        parts.append(chunk.strip())

    for part in parts:
        if not part:
            continue
        key = part.lower()
        if key in MODE_ALIASES:
            mode = MODE_ALIASES[key]
            if mode not in found:
                found.append(mode)
            continue
        for alias, mode in MODE_ALIASES.items():
            if alias in key and mode not in found:
                found.append(mode)
    return found


def sizes_for_equipment(equipment_clean: object, profile: RateCardProfile) -> list[str]:
    if equipment_clean is None or pd.isna(equipment_clean):
        return []
    key = str(equipment_clean).strip()
    sizes = list(EQUIPMENT_SIZE_TARGETS.get(key, []))

    if profile.mirror_40ft_40hc:
        if "40FT" in sizes and "40HC" not in sizes:
            sizes.append("40HC")
        if "40HC" in sizes and "40FT" not in sizes:
            sizes.append("40FT")

    if profile.allowed_sizes is not None:
        allowed = set(profile.allowed_sizes)
        sizes = [s for s in sizes if s in allowed]
        # If equipment is 40' family but mapping missed, still allow profile sizes
        if not sizes and any(token in key for token in ("40'", "40’", "High Cube", "HC")):
            sizes = [s for s in SIZE_ORDER if s in allowed and s != "20FT"]

    return [s for s in SIZE_ORDER if s in sizes]


def modes_for_row(
    origin_mode: object,
    dest_mode: object,
    profile: RateCardProfile,
) -> list[str]:
    if profile.force_targets:
        # Modes implied by forced targets; unused for target building
        return [DEFAULT_MODE]

    if not profile.use_origin_dest_modes:
        return [DEFAULT_MODE]

    values = (dest_mode, origin_mode) if profile.prefer_dest_mode else (origin_mode, dest_mode)
    modes: list[str] = []
    for value in values:
        for mode in parse_modes_from_cell(value):
            if mode in profile.allowed_modes and mode not in modes:
                modes.append(mode)
    if not modes:
        modes = [DEFAULT_MODE] if DEFAULT_MODE in profile.allowed_modes else list(
            profile.allowed_modes[:1]
        )
    return [m for m in MODE_ORDER if m in modes]


def equipment_targets_for_row(
    equipment_clean: object,
    origin_mode: object,
    dest_mode: object,
    profile: RateCardProfile,
) -> list[str]:
    """Mode/size labels that should receive this row's equipment-split costs."""
    if profile.force_targets is not None:
        return list(profile.force_targets)

    sizes = sizes_for_equipment(equipment_clean, profile)
    if not sizes:
        return []
    modes = modes_for_row(origin_mode, dest_mode, profile)
    targets: list[str] = []
    for mode in modes:
        for size in sizes:
            label = f"{mode}/{size}"
            if label not in targets:
                targets.append(label)
    return targets


def origin_mode_series(df: pd.DataFrame) -> pd.Series:
    if "Origin Mode" in df.columns:
        return df["Origin Mode"]
    return pd.Series([pd.NA] * len(df))


def dest_mode_series(df: pd.DataFrame) -> pd.Series:
    for name in ("Dest Mode", "Destination Mode"):
        if name in df.columns:
            return df[name]
    return pd.Series([pd.NA] * len(df))


def collect_equipment_order(df: pd.DataFrame, profile: RateCardProfile) -> list[str]:
    """All Mode/Size combos needed by any lane, in stable mode then size order."""
    if profile.force_targets is not None:
        return list(profile.force_targets)

    origin = origin_mode_series(df)
    dest = dest_mode_series(df)
    needed: set[str] = set()
    for idx, equipment in df["Equipment Type Clean"].items():
        needed.update(
            equipment_targets_for_row(
                equipment, origin.at[idx], dest.at[idx], profile
            )
        )
    order: list[str] = []
    for mode in MODE_ORDER:
        for size in SIZE_ORDER:
            label = f"{mode}/{size}"
            if label in needed:
                order.append(label)
    return order


def has_nonzero_cost(series: pd.Series) -> bool:
    """True if the column has at least one real non-zero cost."""
    numeric = pd.to_numeric(series, errors="coerce")
    return bool(((numeric.notna()) & (numeric != 0)).any())


def select_equipment_cost_columns(
    df: pd.DataFrame, profile: RateCardProfile
) -> list[str]:
    """Equipment-split costs that exist and are not empty/all-zero."""
    if profile.equipment_cost_whitelist is not None:
        candidates = [
            name
            for name in profile.equipment_cost_whitelist
            if name in df.columns and has_nonzero_cost(df[name])
        ]
        # Preserve whitelist order, unique
        seen: set[str] = set()
        ordered: list[str] = []
        for name in candidates:
            if name not in seen:
                seen.add(name)
                ordered.append(name)
        return ordered

    return [
        name
        for name in EQUIPMENT_COST_COLUMNS
        if name in df.columns and has_nonzero_cost(df[name])
    ]


def select_shared_cost_columns(
    df: pd.DataFrame, profile: RateCardProfile
) -> list[str]:
    """Shared end costs — include when present, even if all zeros."""
    if not profile.include_shared_costs:
        return []
    return [name for name in SHARED_COST_COLUMNS if name in df.columns]


def build_equipment_cost_block(
    df: pd.DataFrame,
    cost_name: str,
    equipment: str,
    origin_modes: pd.Series,
    dest_modes: pd.Series,
    profile: RateCardProfile,
) -> CostBlock:
    currency = pd.Series([pd.NA] * len(df), dtype="object")
    values = pd.Series([pd.NA] * len(df), dtype="object")
    for idx, raw_equipment in df["Equipment Type Clean"].items():
        targets = equipment_targets_for_row(
            raw_equipment,
            origin_modes.at[idx],
            dest_modes.at[idx],
            profile,
        )
        if equipment in targets:
            currency.at[idx] = df.at[idx, "Currency"]
            values.at[idx] = df.at[idx, cost_name]
    return CostBlock(
        title=cost_title(cost_name, equipment),
        rate_by=rate_by_label(equipment),
        basis=rate_basis(cost_name),
        currency=currency,
        values=values,
    )


def build_shared_cost_block(df: pd.DataFrame, cost_name: str) -> CostBlock:
    """One cost block for all rows / equipment types."""
    return CostBlock(
        title=cost_title(cost_name),
        rate_by=SHARED_RATE_BY,
        basis=rate_basis(cost_name),
        currency=df["Currency"].copy(),
        values=df[cost_name].copy(),
    )


def build_result_matrix(
    df: pd.DataFrame, source: Path | str | None = None
) -> ResultMatrix:
    """Build shipment details + paired Currency/value cost blocks."""
    required = [
        "Rate ID",
        "Origin Country Clean",
        "Origin City Clean",
        "Destination Country Clean",
        "Destination City Clean",
        "Equipment Type Clean",
        "Service Level Clean",
        "Carrier Name",
        "Rate Effective Date",
        "Rate Expiry Date",
        "Currency",
    ]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise KeyError(f"Missing required columns: {missing}")

    profile = detect_rate_card_profile(source)
    working = filter_profile_rows(df, profile)
    if profile.service_label:
        working = apply_service_label(working, profile.service_label)
    else:
        working = working.copy()

    if working.empty:
        raise ValueError(
            f"No rows left after filters for profile '{profile.name}' "
            f"(source: {Path(source).name if source else 'unknown'})"
        )

    shipment = pd.DataFrame(
        {
            "Rate ID": working["Rate ID"],
            "Transport Mode": optional_series(working, "Transport Mode", "Mode Clean"),
            "Origin Country": working["Origin Country Clean"],
            "Origin City": working["Origin City Clean"],
            "Origin Postal Code": optional_series(working, "Origin Postal Code"),
            "Destination Country": working["Destination Country Clean"],
            "Destination City": working["Destination City Clean"],
            "Destination Postal Code": optional_series(working, "Destination Postal Code"),
            "Business Unit Name": optional_series(
                working, "Business Unit Name", "Business Unit"
            ),
            "equipment_type": working["Equipment Type Clean"],
            "Service": working["Service Level Clean"],
            "Carrier Name": [
                map_carrier_name(c, d)
                for c, d in zip(
                    working["Carrier Name"],
                    working["Destination Country Clean"],
                    strict=True,
                )
            ],
            "Valid from": working["Rate Effective Date"].map(format_date_ddmmyyyy),
            "Valid to": working["Rate Expiry Date"].map(format_date_ddmmyyyy),
        }
    )

    origin_modes = origin_mode_series(working)
    dest_modes = dest_mode_series(working)
    equipment_order = collect_equipment_order(working, profile)
    equipment_costs = select_equipment_cost_columns(working, profile)

    costs: list[CostBlock] = []
    for equipment in equipment_order:
        for cost_name in equipment_costs:
            costs.append(
                build_equipment_cost_block(
                    working,
                    cost_name,
                    equipment,
                    origin_modes,
                    dest_modes,
                    profile,
                )
            )

    for cost_name in select_shared_cost_columns(working, profile):
        costs.append(build_shared_cost_block(working, cost_name))

    return ResultMatrix(shipment=shipment, costs=costs)


DOMESTIC_CARRIER = "DSV ZA"
DOMESTIC_SHEET_NAME = "DSV ZA (Domestic)"


def build_domestic_matrix(result: ResultMatrix) -> ResultMatrix | None:
    """Duplicate DSV ZA shipment lanes with Customs VAT + Finance Fee only."""
    carrier = result.shipment["Carrier Name"].astype(str).str.strip()
    mask = carrier == DOMESTIC_CARRIER
    if not mask.any():
        return None

    shipment = result.shipment.loc[mask].copy().reset_index(drop=True)
    shipment["Carrier Name"] = shipment["Carrier Name"].map(
        lambda name: f"{str(name).strip()} (Domestic)"
    )
    n = len(shipment)

    costs = [
        CostBlock(
            title="Customs VAT",
            rate_by=SHARED_RATE_BY,
            basis="p/unit",
            left_header="Currency",
            currency=pd.Series(["ZAR"] * n, dtype="object"),
            values=pd.Series([1] * n, dtype="object"),
        ),
        CostBlock(
            title="Finance Fee",
            rate_by=SHARED_RATE_BY,
            basis="Over costs",
            left_header="%",
            currency=pd.Series([pd.NA] * n, dtype="object"),
            values=pd.Series([0.4] * n, dtype="object"),
        ),
    ]
    return ResultMatrix(shipment=shipment, costs=costs)


def _style_header_cell(cell) -> None:
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = CENTER
    cell.border = THIN


def _write_result_sheet(ws, result: ResultMatrix) -> None:
    """Write screenshot-like layout: shipment labels on row 3; cost pairs merged above."""
    shipment_cols = list(result.shipment.columns)
    n_ship = len(shipment_cols)
    n_rows = len(result.shipment)
    total_cols = n_ship + 2 * len(result.costs)

    for col_idx, name in enumerate(shipment_cols, start=1):
        for row in (1, 2):
            cell = ws.cell(row, col_idx, None)
            _style_header_cell(cell)
        cell = ws.cell(3, col_idx, name)
        _style_header_cell(cell)

    for i, block in enumerate(result.costs):
        start_col = n_ship + 1 + i * 2
        end_col = start_col + 1

        ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        ws.merge_cells(start_row=2, start_column=start_col, end_row=2, end_column=end_col)

        top = ws.cell(1, start_col, block.title)
        mid = ws.cell(2, start_col, block.rate_by)
        left = ws.cell(3, start_col, block.left_header)
        right = ws.cell(3, end_col, block.basis)

        for cell in (top, mid, left, right, ws.cell(1, end_col), ws.cell(2, end_col)):
            _style_header_cell(cell)

    for r_offset in range(n_rows):
        excel_row = 4 + r_offset
        zebra = r_offset % 2 == 1

        for c_idx, col_name in enumerate(shipment_cols, start=1):
            value = result.shipment.iloc[r_offset, c_idx - 1]
            cell = ws.cell(excel_row, c_idx, None if pd.isna(value) else value)
            cell.alignment = LEFT
            cell.border = THIN
            if zebra:
                cell.fill = ZEBRA_FILL

        for i, block in enumerate(result.costs):
            curr_col = n_ship + 1 + i * 2
            val_col = curr_col + 1
            curr_val = block.currency.iloc[r_offset]
            amt_val = block.values.iloc[r_offset]

            curr_cell = ws.cell(
                excel_row, curr_col, None if pd.isna(curr_val) else curr_val
            )
            amt_cell = ws.cell(excel_row, val_col, None if pd.isna(amt_val) else amt_val)
            for cell in (curr_cell, amt_cell):
                cell.alignment = LEFT
                cell.border = THIN
                if zebra:
                    cell.fill = ZEBRA_FILL

    if total_cols > 0:
        last_col = get_column_letter(total_cols)
        last_row = max(3, 3 + n_rows)
        ws.auto_filter.ref = f"A3:{last_col}{last_row}"

    for col_idx in range(1, total_cols + 1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = 16 if col_idx <= n_ship else 12

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 18
    ws.freeze_panes = None


def save_result_matrix(
    result: ResultMatrix,
    source: Path,
    domestic: ResultMatrix | None = None,
) -> Path:
    """Write main result sheet and optional DSV ZA domestic sheet."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = OUTPUT_DIR / f"{source.stem}_result.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "result"
    _write_result_sheet(ws, result)

    if domestic is not None and len(domestic.shipment) > 0:
        domestic_ws = wb.create_sheet(DOMESTIC_SHEET_NAME)
        _write_result_sheet(domestic_ws, domestic)

    wb.save(output_path)
    return output_path


def main() -> None:
    files = list_processing_files()
    if not files:
        print(f"No supported files found in {PROCESSING_DIR}")
        print("Run process_input.py first to create a processing file.")
        return

    selected = ask_user_to_select(files)
    print(f"\nLoading: {selected.name}")
    df = load_to_dataframe(selected)
    print(f"Loaded DataFrame with shape {df.shape[0]} rows x {df.shape[1]} columns")

    print("Building result matrix...")
    result = build_result_matrix(df, source=selected)
    domestic = build_domestic_matrix(result)
    result_path = save_result_matrix(result, selected, domestic=domestic)
    print(
        f"Saved result matrix to: {result_path} "
        f"({result.shape[0]} rows x {result.shape[1]} columns)"
    )
    if domestic is not None:
        print(
            f"Added sheet '{DOMESTIC_SHEET_NAME}' "
            f"({domestic.shape[0]} rows x {domestic.shape[1]} columns)"
        )


if __name__ == "__main__":
    main()
